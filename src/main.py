
import asyncio
import csv
import json
import os
from utils import Matrix, Entity, Process
from utils import Generator
from utils.bsp import run_bsp, compute_isa_metrics, as_is_to_clusters
import fitz
import pandas as pd
import matplotlib.pyplot as plt
import textwrap

def export_clustered_excel(result, filename):
    df = result.reordered_matrix.fillna("")

    # Re-sort rows and columns by cluster membership so each cluster's
    # processes and entities are contiguous, producing a clean block-diagonal.
    proc_order = [p for c in result.clusters for p in c.processes if p in df.index]
    ent_order  = [e for c in result.clusters for e in c.entities  if e in df.columns]
    proc_order += [p for p in df.index   if p not in proc_order]
    ent_order  += [e for e in df.columns if e not in ent_order]
    df = df.loc[proc_order, ent_order]

    # Create a writer object
    with pd.ExcelWriter(filename, engine='openpyxl') as writer:
        df.to_excel(writer, sheet_name='Clustered Matrix')
        workbook  = writer.book
        worksheet = writer.sheets['Clustered Matrix']
        
        # Define a list of soft colors for clusters
        colors = ["E1F5FE", "F1F8E9", "FFFDE7", "F3E5F5", "E8EAF6", "E0F2F1"]
        
        from openpyxl.styles import PatternFill
        
        # Apply highlights based on cluster boundaries
        for i, cluster in enumerate(result.clusters):
            fill = PatternFill(start_color=colors[i % len(colors)], 
                               end_color=colors[i % len(colors)], 
                               fill_type="solid")
            
            # Find the row/column indices in the reordered dataframe
            row_indices = [df.index.get_loc(p) + 2 for p in cluster.processes] # +2 for header/1-index
            col_indices = [df.columns.get_loc(e) + 2 for e in cluster.entities]
            
            for r in row_indices:
                for c in col_indices:
                    worksheet.cell(row=r, column=c).fill = fill

def display_clustered_matrix(result, title="BSP Clustered Matrix"):
    df = result.reordered_matrix.fillna("")

    proc_order = [p for c in result.clusters for p in c.processes if p in df.index]
    ent_order  = [e for c in result.clusters for e in c.entities  if e in df.columns]
    proc_order += [p for p in df.index   if p not in proc_order]
    ent_order  += [e for e in df.columns if e not in ent_order]
    df = df.loc[proc_order, ent_order]

    colors = ["#E1F5FE", "#F1F8E9", "#FFFDE7", "#F3E5F5", "#E8EAF6", "#E0F2F1"]

    cell_colors = [["white"] * len(df.columns) for _ in df.index]
    for i, cluster in enumerate(result.clusters):
        for p in cluster.processes:
            if p in df.index:
                for e in cluster.entities:
                    if e in df.columns:
                        r = list(df.index).index(p)
                        c = list(df.columns).index(e)
                        cell_colors[r][c] = colors[i % len(colors)]

    wrap_width = 12
    wrapped_cols = ["\n".join(textwrap.wrap(col, width=wrap_width)) for col in df.columns]
    max_lines = max(len(w.split("\n")) for w in wrapped_cols)

    fig_w = max(14, len(df.columns) * 0.8 + 4)
    fig_h = max(6, len(df.index) * 0.4 + max_lines * 0.3 + 1)
    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    ax.set_title(title, fontsize=10, pad=10)
    ax.axis("off")

    table = ax.table(
        cellText=df.values,
        rowLabels=df.index,
        colLabels=wrapped_cols,
        cellColours=cell_colors,
        loc="center"
    )
    table.auto_set_font_size(False)
    table.set_fontsize(8)
    table.auto_set_column_width(col=list(range(-1, len(df.columns))))

    # Taller header row to fit wrapped text
    header_height = (max_lines * 0.3) / fig_h
    for j in range(len(df.columns)):
        table[0, j].set_height(header_height)

    plt.tight_layout()
    plt.show(block=False)
    plt.pause(0.1)

def review_process_types(process_types: dict) -> dict:
    names = list(process_types.keys())
    print("\nProcess classification — this determines cluster boundaries more than anything else:")
    print("  atomic     = single-system ACID transaction; its written entities get force-merged into one cluster")
    print("  end_to_end = spans multiple departments/systems; allowed to span clusters")
    for i, name in enumerate(names, 1):
        print(f"  [{i}] {name}: {process_types[name]}")
    print("\nEnter '<number> <atomic|end_to_end|ambiguous>' to override a process, one per line. Blank line to accept and continue.")
    while True:
        line = input("> ").strip()
        if not line:
            break
        parts = line.split()
        if len(parts) != 2 or not parts[0].isdigit():
            print("  Format: '<number> <atomic|end_to_end|ambiguous>'")
            continue
        idx, new_type = int(parts[0]), parts[1].lower()
        if not (1 <= idx <= len(names)) or new_type not in ("atomic", "end_to_end", "ambiguous"):
            print("  Invalid number or type.")
            continue
        process_types[names[idx - 1]] = new_type
        print(f"  {names[idx - 1]} -> {new_type}")
    return process_types

def pick_constraint(inferred_constraints: list[str]) -> str:
    if inferred_constraints:
        print("\nSelect a constraint:")
        for i, c in enumerate(inferred_constraints, 1):
            print(f"  [{i}] {c}")
        print(f"  [{len(inferred_constraints) + 1}] other")
        while True:
            choice = input("> ").strip()
            if choice.isdigit():
                idx = int(choice)
                if 1 <= idx <= len(inferred_constraints):
                    return inferred_constraints[idx - 1]
                if idx == len(inferred_constraints) + 1:
                    return input("Enter your architectural constraint: ").strip()
            print(f"Please enter a number between 1 and {len(inferred_constraints) + 1}.")
    return input("Enter your architectural constraint: ").strip()

def read_inputs():

    proc_input = input("Enter processes (comma-separated) OR path to .csv: ").strip()
    if proc_input.lower().endswith('.csv') and os.path.exists(proc_input):
        try:
            with open(proc_input, mode='r', encoding='utf-8') as f:
                reader = csv.reader(f)
                processes = [row[0].strip() for row in reader if row and row[0].strip()]
        except UnicodeDecodeError:
            with open(proc_input, mode='r', encoding='cp1252') as f:
                reader = csv.reader(f)
                processes = [row[0].strip() for row in reader if row and row[0].strip()]
    else:
        processes = [p.strip() for p in proc_input.split(',') if p.strip()]

    info_input = input(
        "Enter information text OR path to file (.txt or .pdf) "
        "— leave blank to generate from the process list alone: "
    ).strip()
    
    if os.path.exists(info_input):
        # Handle PDF
        if info_input.lower().endswith('.pdf'):
            text = ""
            with fitz.open(info_input) as doc:
                for page in doc:
                    text += page.get_text()
            context = text
        # Handle TXT
        elif info_input.lower().endswith('.txt'):
            with open(info_input, "r", encoding='utf-8') as f:
                context = f.read()
    else:
        context = info_input

    return processes, context

async def main():

    processes, context = read_inputs()

    ai = Generator()

    print(f"Extracting CRUD matrix ({len(context)} chars, {len(processes)} processes)...")

    final_result = await ai.extract(
            context=context,
            process_list=processes
        )

    with open("full_extraction_log.json", "w") as f:
        json.dump(final_result.model_dump(), f, indent=4)

    print("Full result saved to full_extraction_log.json")

    matrix = Matrix()

    try:
        with open("full_extraction_log.json", "r") as f:
            structured_data = json.load(f)
    except FileNotFoundError:
        print("Error: full_extraction_log.json not found.")
        return
    
    inferred_constraints = [
        c["description"] for c in structured_data.get("inferred_constraints", [])
    ]
    if inferred_constraints:
        print(f"\n{len(inferred_constraints)} architectural constraint(s) inferred from the document:")
        for c in inferred_constraints:
            print(f"  - {c}")

    defined_entities = {e["name"] for e in structured_data.get("entities", [])}

    for p in structured_data.get("processes", []):
        matrix.set_process_type(p["name"], p["process_type"])

    for op in structured_data.get("operations", []):
        if op["entity_name"] not in defined_entities:
            print(f"  WARNING: skipping op — entity '{op['entity_name']}' not in defined list")
            continue
        matrix.add_entry(
            p_name=op["process_name"],
            e_name=op["entity_name"],
            operation=op["operation"]
        )

    entity_names_in_ops = {op["entity_name"] for op in structured_data.get("operations", [])}
    for e in structured_data.get("entities", []):
        if e["name"] not in entity_names_in_ops:
            print(f"  WARNING: entity '{e['name']}' has no operations — it won't appear in the matrix")

    output_csv = "final_crud_matrix.csv"
    matrix.export_to_csv(output_csv)
    
    print(f"Matrix generated! {output_csv} contains only the AI-identified data.")

    matrix.process_types = review_process_types(matrix.process_types)

    print("Running initial BSP Clustering...")
    initial_bsp = run_bsp(matrix.matrix, process_types=matrix.process_types)

    export_clustered_excel(initial_bsp, "initial_clustered_matrix.xlsx")
    print("Clustered Excel generated: initial_clustered_matrix.xlsx")
    display_clustered_matrix(initial_bsp, title="Initial BSP Clustering")

    with open("resources/ea_principles.txt", "r") as f:
        baseline_principles = f.read()

    accumulated_user_constraints = ""
    accumulated_weights = {}
    accumulated_reasoning = {}
    current_bsp = initial_bsp
    iteration = 0
    bsp_iterations = []

    while True:
        answer = input("\nAre you satisfied with the clustering? (yes/no): ").strip().lower()
        if answer in ("yes", "y"):
            break

        constraints = pick_constraint(inferred_constraints)
        if constraints:
            iteration += 1
            accumulated_user_constraints += f"\n\n[Iteration {iteration}]:\n{constraints}"

        print("\nAnalysing constraints and computing EA weights...")
        new_weights, new_reasoning = await ai._compute_ea_weights(
            current_bsp.clusters,
            baseline_principles,
            accumulated_user_constraints
        )
        accumulated_weights.update(new_weights)
        accumulated_reasoning.update(new_reasoning)

        current_bsp = run_bsp(matrix.matrix, ea_weights=accumulated_weights, process_types=matrix.process_types)

        out_file = f"clustered_matrix_iter{iteration}.xlsx"
        export_clustered_excel(current_bsp, out_file)
        print(f"Clustered Excel generated: {out_file}")
        display_clustered_matrix(current_bsp, title=f"BSP Iteration {iteration}")

        bsp_iterations.append({
            "iteration": iteration,
            "user_constraint": constraints,
            "ea_weights": {
                str(k): {"weight": v, "reasoning": new_reasoning.get(k, "")}
                for k, v in new_weights.items()
            },
            "clusters": [c.to_dict() for c in current_bsp.clusters],
            "entity_owners": current_bsp.entity_owners,
        })

    final_bsp = current_bsp

    tobe_metrics = compute_isa_metrics(final_bsp.clusters, matrix.matrix, matrix.process_types)

    # Build a human-readable weight reasoning trail for the LLM
    weight_reasoning_lines = []
    for k, reasoning in accumulated_reasoning.items():
        weight = accumulated_weights.get(k, 0)
        weight_reasoning_lines.append(f"  {k[0]} <-> {k[1]}: {weight:+.2f} — {reasoning}")
    weight_reasoning_history = "\n".join(weight_reasoning_lines)

    print("\nGenerating system descriptions and market comparisons...")
    systems_analysis = await ai.describe_systems(
        final_clusters=final_bsp.clusters,
        user_constraints_history=accumulated_user_constraints,
        weight_reasoning_history=weight_reasoning_history,
    )

    print("\n" + "=" * 60)
    print("FINAL SYSTEMS ANALYSIS")
    print("=" * 60)
    for s in systems_analysis.systems:
        print(f"\nCluster {s.cluster_id}: {s.suggested_name}")
        print(f"  {s.description}")
        print(f"  Recommendation: {s.build_or_buy}")
        print(f"  {s.build_or_buy_rationale}")
        if s.market_options:
            print("  Market options:")
            for opt in s.market_options:
                print(f"    - {opt.name}: {opt.fit_rationale}")
    print("=" * 60)

    print("\nEvaluating EA principles compliance...")
    ea_compliance = await ai.evaluate_ea_compliance(
        final_clusters=final_bsp.clusters,
        ea_principles=baseline_principles,
        weight_reasoning_history=weight_reasoning_history,
    )

    print("\n" + "=" * 60)
    print("EA PRINCIPLES COMPLIANCE")
    print("=" * 60)
    for ev in ea_compliance.evaluations:
        status_label = {"compliant": "OK", "partial": "PARTIAL", "violated": "VIOLATED"}[ev.status]
        print(f"\n[{status_label}] {ev.principle}")
        print(f"  {ev.justification}")
    print("=" * 60)

    # ── As-Is vs. To-Be comparison (optional) ──────────────────────────
    comparison_result = None
    do_compare = input("\nDo you want to compare with your current application landscape? (yes/no): ").strip().lower()
    if do_compare in ("yes", "y"):
        as_is_path = input("Path to As-Is CSV (two columns: Process, Application): ").strip()
        if os.path.exists(as_is_path):
            as_is_mapping: dict = {}
            with open(as_is_path, mode="r", encoding="utf-8") as f:
                reader = csv.reader(f)
                for row in reader:
                    if len(row) >= 2:
                        proc, app = row[0].strip(), row[1].strip()
                        if proc and app:
                            as_is_mapping.setdefault(app, []).append(proc)

            asis_clusters = as_is_to_clusters(as_is_mapping, matrix.matrix)
            asis_metrics  = compute_isa_metrics(asis_clusters, matrix.matrix, matrix.process_types)

            print("\n" + "=" * 60)
            print("ISA METRICS — AS-IS vs. TO-BE")
            print("=" * 60)
            metric_explanations = {
                "RSF":    "avg IS blocks per process         (ideal 1.0 = min coupling)",
                "NAIEF":  "entity write authority            (ideal 1.0 = single source of truth)",
                "LCOISF": "cluster cohesion                  (ideal 1.0 = no giant-stain clusters)",
                "CPSMF":  "critical/non-critical isolation   (ideal 1.0 = perfect separation)",
            }
            print(f"\n{'Metric':<8} {'As-Is':>6} {'To-Be':>6} {'Δ':>7}   Meaning")
            print("-" * 70)
            for m, explanation in metric_explanations.items():
                asis_val  = asis_metrics.get(m, 0.0)
                tobe_val  = tobe_metrics.get(m, 0.0)
                delta     = tobe_val - asis_val
                arrow     = "↑" if delta > 0 else ("↓" if delta < 0 else "=")
                print(f"{m:<8} {asis_val:>6.3f} {tobe_val:>6.3f} {delta:>+6.3f} {arrow}   {explanation}")
            print()
            print("What these metrics mean:")
            for m, explanation in metric_explanations.items():
                print(f"  {m}: {explanation}")

            print("\nRunning As-Is vs. To-Be analysis...")
            comparison_result = await ai.compare_as_is_to_be(
                as_is_mapping=as_is_mapping,
                final_clusters=final_bsp.clusters,
                systems_analysis=systems_analysis,
                asis_metrics=asis_metrics,
                tobe_metrics=tobe_metrics,
            )

            print("\n" + "=" * 60)
            print("AS-IS vs. TO-BE GAP ANALYSIS")
            print("=" * 60)
            print(f"\n{comparison_result.overall_summary}")
            print(f"\nEstimated complexity: {comparison_result.estimated_complexity}")
            print(f"  {comparison_result.complexity_rationale}")

            print("\n--- Alignment by proposed system ---")
            for a in comparison_result.alignments:
                print(f"\n  [{a.proposed_cluster_id}] {a.proposed_system_name}")
                print(f"    Current apps : {', '.join(a.current_applications) or '—'}")
                if a.processes_to_acquire:
                    print(f"    Move IN      : {', '.join(a.processes_to_acquire)}")
                if a.processes_to_release:
                    print(f"    Move OUT     : {', '.join(a.processes_to_release)}")
                print(f"    {a.alignment_summary}")

            print("\n--- Migration steps ---")
            for step in comparison_result.transformation_steps:
                print(f"\n  Step {step.step_number} [{step.action.upper()}]")
                print(f"    {step.description}")
                print(f"    Why: {step.rationale}")
            print("=" * 60)
        else:
            print(f"  File not found: {as_is_path}. Skipping comparison.")

    with open("full_extraction_log.json", "w") as f:
        json.dump({
            "extraction": final_result.model_dump(),
            "initial_bsp": {
                "clusters": [c.to_dict() for c in initial_bsp.clusters],
                "entity_owners": initial_bsp.entity_owners,
            },
            "bsp_iterations": bsp_iterations,
            "final_bsp": {
                "clusters": [c.to_dict() for c in final_bsp.clusters],
                "entity_owners": final_bsp.entity_owners,
                "ea_weights": {
                    str(k): {"weight": v, "reasoning": accumulated_reasoning.get(k, "")}
                    for k, v in accumulated_weights.items()
                },
                "user_constraints": accumulated_user_constraints.strip(),
            },
            "systems_analysis": [
                {
                    "cluster_id": s.cluster_id,
                    "suggested_name": s.suggested_name,
                    "description": s.description,
                    "build_or_buy": s.build_or_buy,
                    "build_or_buy_rationale": s.build_or_buy_rationale,
                    "market_options": [
                        {"name": o.name, "fit_rationale": o.fit_rationale}
                        for o in s.market_options
                    ],
                }
                for s in systems_analysis.systems
            ],
            "ea_compliance": [
                {
                    "principle": ev.principle,
                    "status": ev.status,
                    "justification": ev.justification,
                }
                for ev in ea_compliance.evaluations
            ],
            "isa_metrics_tobe": tobe_metrics,
            "as_is_comparison": {
                "overall_summary": comparison_result.overall_summary,
                "estimated_complexity": comparison_result.estimated_complexity,
                "complexity_rationale": comparison_result.complexity_rationale,
                "alignments": [
                    {
                        "proposed_cluster_id": a.proposed_cluster_id,
                        "proposed_system_name": a.proposed_system_name,
                        "current_applications": a.current_applications,
                        "processes_to_acquire": a.processes_to_acquire,
                        "processes_to_release": a.processes_to_release,
                        "alignment_summary": a.alignment_summary,
                    }
                    for a in comparison_result.alignments
                ],
                "transformation_steps": [
                    {
                        "step_number": s.step_number,
                        "action": s.action,
                        "description": s.description,
                        "rationale": s.rationale,
                        "affected_processes": s.affected_processes,
                        "affected_applications": s.affected_applications,
                    }
                    for s in comparison_result.transformation_steps
                ],
            } if comparison_result else None,
        }, f, indent=4)

    print("Full extraction log updated with clusters and systems analysis!")

if __name__ == "__main__":
    asyncio.run(main())